import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from core import Store, read_document, export_word, export_slides, export_table, ask_ai, unique_path, protect, OPENROUTER_DEFAULTS, parse_provider


class CoreTests(unittest.TestCase):
    def test_openrouter_defaults_key_and_provider_request(self):
        self.store.save_profile({'agency': 'Trường kiểm thử', 'model': 'old-model', 'provider_routing': 'old-provider'})
        self.store.save_key('fake-openrouter-key-for-test')
        self.store.save_profile({'provider': 'OpenRouter'})
        profile = self.store.profile()
        self.assertEqual(profile['agency'], 'Trường kiểm thử')
        for name, value in OPENROUTER_DEFAULTS.items():
            self.assertEqual(profile[name], value)
        self.assertNotIn('fake-openrouter-key-for-test', (self.store.root / 'api-key.dpapi').read_text())
        self.assertNotIn(b'fake-openrouter-key-for-test', (self.store.root / 'du-lieu.sqlite3').read_bytes())
        seen = {}
        class Response:
            status_code = 200
            def json(self): return {'choices': [{'message': {'content': 'OK'}}]}
        def post(url, **kwargs):
            seen.update(url=url, **kwargs)
            return Response()
        self.assertEqual(ask_ai(profile, self.store.key(), [{'role': 'user', 'content': 'Test'}], post=post), 'OK')
        self.assertEqual(seen['json']['model'], OPENROUTER_DEFAULTS['model'])
        if OPENROUTER_DEFAULTS.get('provider_routing'):
            self.assertEqual(seen['json']['provider'], parse_provider(OPENROUTER_DEFAULTS['provider_routing']))
        else:
            self.assertNotIn('provider', seen['json'])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = Store(self.root / 'data')

    def tearDown(self):
        self.store.db.close()
        self.tmp.cleanup()

    def test_persistence_and_backup_excludes_key(self):
        self.store.put('task', {'title': 'Nộp báo cáo tiếng Việt', 'due': '2026-09-30'})
        self.store.save_key('test-secret-only')
        self.assertEqual(self.store.key(), 'test-secret-only')
        content = self.store.root / 'bao-cao' / 'Báo cáo.txt'
        content.write_text('Số liệu: 128', encoding='utf-8')
        snapshot = self.root / 'backup.zip'
        self.store.backup(snapshot)
        with zipfile.ZipFile(snapshot) as z:
            self.assertNotIn('api-key.dpapi', z.namelist())
            z.extractall(self.root / 'restore')
        restored = Store(self.root / 'restore')
        try:
            self.assertEqual(restored.all('task')[0]['title'], 'Nộp báo cáo tiếng Việt')
            self.assertEqual((restored.root / 'bao-cao' / 'Báo cáo.txt').read_text(encoding='utf-8'), 'Số liệu: 128')
        finally:
            restored.db.close()

    def test_office_roundtrip_and_no_overwrite(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import load_workbook
        doc = self.root / 'văn bản.docx'
        export_word(doc, 'Báo cáo', 'Số học sinh: 128\nNội dung cần bổ sung', {'agency': 'Sở Giáo dục'})
        self.assertIn('Số học sinh: 128', read_document(doc))
        self.assertTrue(Document(doc).tables)
        next_path = unique_path(self.root, doc.name)
        self.assertEqual(next_path.name, 'văn bản-v2.docx')
        next_path.touch()
        self.assertEqual(unique_path(self.root, doc.name).name, 'văn bản-v3.docx')
        slides = self.root / 'slides.pptx'
        export_slides(slides, 'Báo cáo', '\n'.join('Dòng số ' + str(i) for i in range(22)))
        presentation = Presentation(slides)
        text = '\n'.join(s.text for slide in presentation.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('Dòng số 21', text)
        self.assertGreater(len(presentation.slides), 2)
        xlsx = self.root / 'table.xlsx'
        export_table(xlsx, [['=HYPERLINK("https://example.invalid")', 'Tiếng Việt']], ['Công việc', 'Ghi chú'])
        wb = load_workbook(xlsx)
        self.assertEqual(wb.active['A2'].data_type, 's')
        self.assertEqual(wb.active['B2'].value, 'Tiếng Việt')
        wb.close()
        self.assertIn('Tiếng Việt', read_document(xlsx))

    def test_compatible_api_payload_history_and_errors(self):
        seen = []
        class Response:
            status_code = 200
            def json(self):
                return {'choices': [{'message': {'content': 'Đã nhận đủ nội dung.'}}]}
        def post(url, **kwargs):
            seen.append((url, kwargs))
            return Response()
        profile = {'provider': 'API tương thích', 'endpoint': 'http://localhost:1234/v1/', 'model': 'local-model', 'agency': 'Cơ quan A'}
        messages = [{'role': 'user', 'content': 'Soạn báo cáo'}, {'role': 'assistant', 'content': 'Dự thảo'}, {'role': 'user', 'content': 'Sửa phần 2'}]
        self.assertEqual(ask_ai(profile, 'test-only', messages, '128 học sinh', post=post), 'Đã nhận đủ nội dung.')
        url, request = seen[0]
        self.assertEqual(url, 'http://localhost:1234/v1/chat/completions')
        self.assertEqual(request['json']['messages'][1:], messages)
        self.assertIn('128 học sinh', request['json']['messages'][0]['content'])
        self.assertEqual(request['headers']['Authorization'], 'Bearer test-only')
        self.assertFalse(request['allow_redirects'])
        Response.status_code = 401
        with self.assertRaisesRegex(ValueError, '401'):
            ask_ai(profile, 'test-only', messages, post=post)
        with self.assertRaisesRegex(ValueError, 'HTTPS'):
            ask_ai(profile | {'endpoint': 'http://example.com/v1'}, 'test-only', messages, post=post)

    def test_missing_model_never_calls_network(self):
        def forbidden(*args, **kwargs):
            self.fail('No request should be sent')
        with self.assertRaisesRegex(ValueError, 'mô hình'):
            ask_ai({'provider': 'API tương thích'}, '', [], post=forbidden)

    def test_gemini_response_shape(self):
        class Response:
            status_code = 200
            def __init__(self, value): self.value = value
            def json(self): return self.value
        self.assertEqual(ask_ai({'provider': 'Gemini', 'model': 'test-model'}, 'test', [{'role': 'user', 'content': 'Xin chào'}], post=lambda *a, **kw: Response({'candidates': [{'content': {'parts': [{'text': 'Nội bộ', 'thought': True}, {'text': 'Xin chào'}]}}]})), 'Xin chào')


if __name__ == '__main__':
    unittest.main(verbosity=2)
