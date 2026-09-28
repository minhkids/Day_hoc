import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_gradebook_xlsx(data, profile, output_path):
    """
    Tạo Bảng theo dõi điểm số và đánh giá học sinh chuẩn file Excel (.xlsx)
    có định dạng ô màu, kẻ khung và công thức tính toán.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bảng Điểm Học Kỳ"

    # Định dạng Font & Style
    font_header = Font(name="Times New Roman", size=14, bold=True, color="002B6B")
    font_title = Font(name="Times New Roman", size=16, bold=True, color="0B4FC4")
    font_col_header = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
    font_regular = Font(name="Times New Roman", size=12)
    font_bold = Font(name="Times New Roman", size=12, bold=True)

    fill_header = PatternFill(start_color="1462E6", end_color="1462E6", fill_type="solid")
    fill_summary = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='A0B0C8'),
        right=Side(style='thin', color='A0B0C8'),
        top=Side(style='thin', color='A0B0C8'),
        bottom=Side(style='thin', color='A0B0C8')
    )

    # 1. Tiêu đề trường & Bảng điểm
    ws.merge_cells("A1:H1")
    ws["A1"] = profile.get("school_name", "TRƯỜNG THCS").upper()
    ws["A1"].font = font_header
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A3:H3")
    ws["A3"] = f"BẢNG ĐIỂM ĐÁNH GIÁ MÔN {profile.get('subject', 'TOÁN HỌC').upper()}"
    ws["A3"].font = font_title
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A4:H4")
    ws["A4"] = f"Lớp: {data.get('class_name', '7A1')} • Giáo viên bộ môn: {profile.get('full_name', '')} • Năm học: 2026 - 2027"
    ws["A4"].font = Font(name="Times New Roman", size=12, italic=True)
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    # 2. Hàng tiêu đề cột
    headers = ["STT", "Họ và tên học sinh", "Điểm TX 1", "Điểm TX 2", "Điểm TX 3", "Điểm GK", "Điểm CK", "ĐTB Môn"]
    header_row = 6

    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=text)
        cell.font = font_col_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # 3. Dữ liệu học sinh mẫu
    students = data.get("students", [
        ("Nguyễn Hoàng An", 8.5, 9.0, 8.0, 8.5, 9.0),
        ("Trần Thị Bình", 7.0, 7.5, 8.0, 7.0, 7.5),
        ("Lê Văn Cường", 9.0, 9.5, 9.0, 9.0, 9.5),
        ("Phạm Thị Dung", 6.5, 7.0, 6.0, 6.5, 7.0),
        ("Hoàng Minh Đức", 8.0, 8.5, 7.5, 8.0, 8.5),
        ("Đỗ Ngọc Hạnh", 9.5, 10.0, 9.5, 10.0, 9.5),
        ("Bùi Quốc Huy", 5.5, 6.0, 6.5, 5.0, 6.0),
        ("Vũ Phương Mai", 8.0, 8.0, 8.5, 8.5, 8.0),
        ("Đặng Văn Nam", 7.5, 8.0, 7.0, 7.5, 8.0),
        ("Ngô Thu Thảo", 9.0, 8.5, 9.0, 9.5, 9.0),
    ])

    start_row = 7
    for idx, stu in enumerate(students, start=1):
        r = start_row + idx - 1
        name, tx1, tx2, tx3, gk, ck = stu

        ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=name).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=3, value=tx1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=tx2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=tx3).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=6, value=gk).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=7, value=ck).alignment = Alignment(horizontal="center")

        # Công thức tính ĐTB Môn: (TX1 + TX2 + TX3 + GK*2 + CK*3) / 8
        formula = f"=ROUND((C{r} + D{r} + E{r} + F{r}*2 + G{r}*3)/8, 1)"
        cell_dtb = ws.cell(row=r, column=8, value=formula)
        cell_dtb.alignment = Alignment(horizontal="center")
        cell_dtb.font = font_bold

        # Kẻ khung
        for col_idx in range(1, 9):
            ws.cell(row=r, column=col_idx).border = thin_border
            if col_idx in [1, 2]:
                ws.cell(row=r, column=col_idx).font = font_regular

    # 4. Hàng tổng kết điểm trung bình cả lớp
    end_data_row = start_row + len(students) - 1
    summary_row = end_data_row + 1

    ws.merge_cells(f"A{summary_row}:B{summary_row}")
    cell_sum_label = ws.cell(row=summary_row, column=1, value="ĐIỂM TRUNG BÌNH CẢ LỚP:")
    cell_sum_label.font = font_bold
    cell_sum_label.alignment = Alignment(horizontal="center", vertical="center")

    for c in ["C", "D", "E", "F", "G", "H"]:
        col_num = ord(c) - ord('A') + 1
        cell_avg = ws.cell(row=summary_row, column=col_num, value=f"=ROUND(AVERAGE({c}{start_row}:{c}{end_data_row}), 1)")
        cell_avg.font = font_bold
        cell_avg.alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, 9):
        c_item = ws.cell(row=summary_row, column=col_idx)
        c_item.fill = fill_summary
        c_item.border = thin_border

    # Tự động căn chỉnh độ rộng cột
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len and not cell.coordinate in ["A1", "A3", "A4"]:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)
    ws.column_dimensions['B'].width = 28 # Cột họ tên rộng hơn

    wb.save(output_path)
    return output_path
