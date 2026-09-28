import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Thiết lập padding cho ô bảng trong word"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """Thiết lập đường viền cho ô trong word"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            border = OxmlElement(f'w:{edge}')
            border.set(qn('w:val'), edge_data.get('val', 'single'))
            border.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(border)
    tcPr.append(tcBorders)

def apply_vietnamese_page_setup(doc):
    """Thiết lập lề trang chuẩn Nghị định 30/2020: Trái 30mm, Phải 15mm, Trên 20mm, Dưới 20mm"""
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm
        section.right_margin = Inches(0.59)   # 1.5 cm

def export_lesson_plan_docx(lesson_data, profile, output_path):
    """
    Xuất Kế hoạch bài dạy chuẩn Công văn 5512/BGDĐT sang file .docx
    """
    doc = Document()
    apply_vietnamese_page_setup(doc)

    # Style chuẩn Times New Roman
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0, 0, 0)

    # Header: Đơn vị & Người soạn
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    c_left = table.cell(0, 0)
    p_left = c_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.paragraph_format.line_spacing = 1.15
    run = p_left.add_run(f"{profile.get('school_name', 'TRƯỜNG THCS').upper()}\n")
    run.font.bold = True
    p_left.add_run(f"Tổ: {profile.get('subject', 'Chuyên môn')}")

    c_right = table.cell(0, 1)
    p_right = c_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.paragraph_format.line_spacing = 1.15
    run_r = p_right.add_run("KẾ HOẠCH BÀI DẠY\n")
    run_r.font.bold = True
    p_right.add_run(f"Giáo viên: {profile.get('full_name', 'Giáo viên')}")

    doc.add_paragraph() # Dòng trống

    # Tên bài dạy
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.3
    run_sub = p_title.add_run(f"BÀI HỌC: {lesson_data.get('topic', 'TÊN BÀI DẠY').upper()}\n")
    run_sub.font.bold = True
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = RGBColor(11, 79, 196)

    run_meta = p_title.add_run(f"Môn: {profile.get('subject', '')} - {lesson_data.get('grade', 'Lớp 7')}; Thời lượng: {lesson_data.get('duration', '1 tiết')}")
    run_meta.font.italic = True
    run_meta.font.size = Pt(12)

    # I. YÊU CẦU CẦN ĐẠT
    p_muc1 = doc.add_paragraph()
    p_muc1.paragraph_format.space_before = Pt(8)
    p_muc1.paragraph_format.space_after = Pt(4)
    run_m1 = p_muc1.add_run("I. YÊU CẦU CẦN ĐẠT:")
    run_m1.font.bold = True
    run_m1.font.size = Pt(13.5)

    yccd = lesson_data.get('yccd', {})
    doc.add_paragraph(f"1. Về năng lực đặc thù: {yccd.get('nang_luc_dac_thu', 'Nắm vững kiến thức trọng tâm của bài, vận dụng giải quyết bài toán thực tế.')}")
    doc.add_paragraph(f"2. Về năng lực chung: {yccd.get('nang_luc_chung', 'Tự chủ và tự học; giao tiếp và hợp tác nhóm; giải quyết vấn đề sáng tạo.')}")
    doc.add_paragraph(f"3. Về phẩm chất: {yccd.get('pham_chat', 'Chăm chỉ, trung thực, có trách nhiệm trong học tập và hợp tác nhóm.')}")

    # II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU
    p_muc2 = doc.add_paragraph()
    p_muc2.paragraph_format.space_before = Pt(8)
    p_muc2.paragraph_format.space_after = Pt(4)
    run_m2 = p_muc2.add_run("II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU:")
    run_m2.font.bold = True
    run_m2.font.size = Pt(13.5)

    doc.add_paragraph(f"1. Giáo viên: {lesson_data.get('tb_gv', 'SGK, kế hoạch bài dạy, bài trình chiếu PowerPoint, phiếu học tập, đồ dùng trực quan.')}")
    doc.add_paragraph(f"2. Học sinh: {lesson_data.get('tb_hs', 'SGK, vở ghi, dụng cụ học tập theo yêu cầu của giáo viên.')}")

    # III. TIẾN TRÌNH DẠY HỌC
    p_muc3 = doc.add_paragraph()
    p_muc3.paragraph_format.space_before = Pt(8)
    p_muc3.paragraph_format.space_after = Pt(4)
    run_m3 = p_muc3.add_run("III. TIẾN TRÌNH DẠY HỌC (THEO CÔNG VĂN 5512):")
    run_m3.font.bold = True
    run_m3.font.size = Pt(13.5)

    activities = lesson_data.get('activities', [
        ("Hoạt động 1: Mở đầu / Khởi động (5 phút)", "Tạo hứng thú, kết nối kiến thức cũ với tình huống thực tiễn mở đầu bài mới.", "Trò chơi câu đố / Tình huống khởi động.", "Học sinh tham gia tích cực, nhận diện được vấn đề cần học."),
        ("Hoạt động 2: Hình thành kiến thức mới (18 phút)", "Khám phá kiến thức trọng tâm bài học, xây dựng khái niệm / định lý.", "Tổ chức hoạt động nhóm tìm hiểu tài liệu và thảo luận.", "Học sinh phát biểu và ghi nhớ được quy tắc, nội dung chính."),
        ("Hoạt động 3: Luyện tập (15 phút)", "Củng cố và rèn luyện kỹ năng thực hành bài tập cơ bản.", "Giao bài tập vận dụng trong SGK, gọi học sinh lên bảng trình bày.", "Học sinh hoàn thành bài tập chính xác theo đáp án mẫu."),
        ("Hoạt động 4: Vận dụng (7 phút)", "Mở rộng liên hệ thực tế đời sống, hướng dẫn tự học ở nhà.", "Nêu tình huống ứng dụng thực tế để học sinh liên hệ tìm tòi.", "Học sinh hình thành tư duy vận dụng kiến thức vào cuộc sống.")
    ])

    for title, muc_tieu, noi_dung, san_pham in activities:
        p_act = doc.add_paragraph()
        p_act.paragraph_format.space_before = Pt(6)
        r_act = p_act.add_run(f"● {title}")
        r_act.font.bold = True
        r_act.font.color.rgb = RGBColor(20, 98, 230)

        doc.add_paragraph(f"  a) Mục tiêu: {muc_tieu}")
        doc.add_paragraph(f"  b) Nội dung: {noi_dung}")
        doc.add_paragraph(f"  c) Sản phẩm: {san_pham}")
        doc.add_paragraph("  d) Tổ chức thực hiện: GV chuyển giao nhiệm vụ -> HS tiếp nhận & thực hiện -> Báo cáo thảo luận -> GV đánh giá, chuẩn hóa kết luận.")

    # Chữ ký duyệt cuối trang
    doc.add_paragraph()
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell_duyet = sig_table.cell(0, 0)
    p_duyet = cell_duyet.paragraphs[0]
    p_duyet.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_d = p_duyet.add_run("DUYỆT CỦA TỔ CHUYÊN MÔN\n\n\n\n")
    r_d.font.bold = True
    p_duyet.add_run("(Ký và ghi rõ họ tên)")

    cell_gv = sig_table.cell(0, 1)
    p_gv = cell_gv.paragraphs[0]
    p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_gv.add_run(f"{profile.get('district', '')}, ngày ... tháng ... năm 2026\n")
    r_gv = p_gv.add_run("GIÁO VIÊN SOẠN\n\n\n\n")
    r_gv.font.bold = True
    p_gv.add_run(f"{profile.get('full_name', '')}")

    doc.save(output_path)
    return output_path

def export_admin_doc_docx(doc_data, profile, output_path):
    """
    Xuất Văn bản hành chính (Quyết định / Kế hoạch / Tờ trình) chuẩn Nghị định 30/2020/NĐ-CP
    """
    doc = Document()
    apply_vietnamese_page_setup(doc)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)

    # 1. Header 2 cột: Cơ quan ban hành & Quốc hiệu
    head_table = doc.add_table(rows=1, cols=2)
    head_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    head_table.autofit = True

    # Cột trái
    c_left = head_table.cell(0, 0)
    p_left = c_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.paragraph_format.line_spacing = 1.1
    p_left.add_run(f"UBND {profile.get('district', 'HUYỆN').upper()}\n").font.size = Pt(12)
    run_s = p_left.add_run(f"{profile.get('school_name', 'TRƯỜNG TIỂU HỌC').upper()}\n")
    run_s.font.bold = True
    run_s.font.size = Pt(12.5)
    p_left.add_run(f"Số: {doc_data.get('doc_number', '01')}/{doc_data.get('doc_type_code', 'QĐ')}-{doc_data.get('short_name', 'TH')}").font.size = Pt(12)

    # Cột phải
    c_right = head_table.cell(0, 1)
    p_right = c_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.paragraph_format.line_spacing = 1.1
    run_qh = p_right.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    run_qh.font.bold = True
    run_qh.font.size = Pt(12)
    run_tn = p_right.add_run("Độc lập - Tự do - Hạnh phúc\n")
    run_tn.font.bold = True
    run_tn.font.size = Pt(13)
    run_date = p_right.add_run(f"{profile.get('province', 'Hà Nội')}, ngày ... tháng ... năm 2026")
    run_date.font.italic = True
    run_date.font.size = Pt(12)

    doc.add_paragraph()

    # 2. Tên loại văn bản & Trích yếu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(14)
    run_type = p_title.add_run(f"{doc_data.get('doc_type', 'QUYẾT ĐỊNH')}\n")
    run_type.font.bold = True
    run_type.font.size = Pt(15)

    run_desc = p_title.add_run(f"{doc_data.get('title', 'Về việc ban hành quy chế chuyên môn')}")
    run_desc.font.size = Pt(13)

    # 3. Thẩm quyền ban hành
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_a = p_auth.add_run(f"{profile.get('role', 'HIỆU TRƯỞNG').upper()} {profile.get('school_name', '').upper()}")
    run_a.font.bold = True

    # Căn cứ pháp lý
    cancu_list = doc_data.get('cancu', [
        "Căn cứ Điều lệ trường học ban hành theo Thông tư của Bộ Giáo dục và Đào tạo;",
        "Căn cứ Nghị định số 30/2020/NĐ-CP ngày 05/3/2020 của Chính phủ về công tác văn thư;",
        "Xét đề nghị của Hội đồng thi đua và Tổ trưởng chuyên môn nhà trường."
    ])
    for cc in cancu_list:
        p_cc = doc.add_paragraph()
        p_cc.paragraph_format.line_spacing = 1.2
        r_cc = p_cc.add_run(f"    {cc}")
        r_cc.font.italic = True

    # Lời ban hành
    p_banhanh = doc.add_paragraph()
    p_banhanh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_banhanh.paragraph_format.space_before = Pt(6)
    p_banhanh.paragraph_format.space_after = Pt(6)
    r_bh = p_banhanh.add_run("QUYẾT ĐỊNH:")
    r_bh.font.bold = True

    # Các điều khoản
    articles = doc_data.get('articles', [
        ("Điều 1", "Ban hành kèm theo Quyết định này Kế hoạch tổ chức các hoạt động giáo dục năm học 2026 - 2027."),
        ("Điều 2", "Giao các tổ chuyên môn, giáo viên phụ trách triển khai thực hiện nghiêm túc các nội dung kế hoạch đã đề ra."),
        ("Điều 3", "Các bộ phận liên quan và các cá nhân có tên tại Điều 1 chịu trách nhiệm thi hành Quyết định này./.")
    ])

    for dieu, nd in articles:
        p_dieu = doc.add_paragraph()
        p_dieu.paragraph_format.line_spacing = 1.3
        p_dieu.paragraph_format.space_after = Pt(4)
        r_dieu = p_dieu.add_run(f"    {dieu}. ")
        r_dieu.font.bold = True
        p_dieu.add_run(nd)

    # 4. Nơi nhận & Chữ ký
    doc.add_paragraph()
    foot_table = doc.add_table(rows=1, cols=2)
    foot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    foot_table.autofit = True

    c_rec = foot_table.cell(0, 0)
    p_rec = c_rec.paragraphs[0]
    p_rec.paragraph_format.line_spacing = 1.1
    r_nn = p_rec.add_run("Nơi nhận:\n")
    r_nn.font.bold = True
    r_nn.font.italic = True
    r_nn.font.size = Pt(11)
    p_rec.add_run("- Phòng GD&ĐT (để b/c);\n- Như Điều 3 (để t/h);\n- Lưu: VT.").font.size = Pt(11)

    c_sign = foot_table.cell(0, 1)
    p_sign = c_sign.paragraphs[0]
    p_sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sign.add_run(f"{profile.get('role', 'HIỆU TRƯỞNG').upper()}\n\n\n\n\n")
    r_s.font.bold = True
    r_name = p_sign.add_run(f"{profile.get('full_name', '')}")
    r_name.font.bold = True

    doc.save(output_path)
    return output_path
