import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def export_lesson_slides_pptx(slide_data, profile, output_path):
    """
    Tạo bộ Slide PowerPoint bài giảng (.pptx) chuẩn 16:9 với màu sắc hiện đại,
    kèm ghi chú thuyết minh (Speaker Notes) dưới từng slide mà không cần cài đặt MS PowerPoint.
    """
    prs = Presentation()
    # Kích thước chuẩn 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Layout trống tự do tùy biến

    # Bảng màu
    PRIMARY_COLOR = RGBColor(20, 98, 230)    # Royal Blue
    NAVY_COLOR = RGBColor(11, 43, 107)       # Dark Navy
    TEXT_MAIN = RGBColor(42, 54, 79)         # Dark Gray Text
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(244, 248, 254)

    # 1. Slide 1: Trang bìa (Title Slide)
    s1 = prs.slides.add_slide(blank_layout)

    # Background khối màu
    top_shape = s1.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(4.8)) # MSO_SHAPE.RECTANGLE
    top_shape.fill.solid()
    top_shape.fill.fore_color.rgb = NAVY_COLOR
    top_shape.line.color.rgb = NAVY_COLOR

    # Tiêu đề bài dạy
    tx_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(2.5))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = f"{profile.get('school_name', '').upper()} • {profile.get('subject', '').upper()}"
    p_badge.font.size = Pt(16)
    p_badge.font.color.rgb = RGBColor(200, 220, 255)
    p_badge.font.bold = True

    p_title = tf.add_paragraph()
    p_title.text = slide_data.get('topic', 'BÀI GIẢNG ĐIỆN TỬ').upper()
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE

    # Thông tin giáo viên bên dưới
    tx_info = s1.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.333), Inches(1.8))
    tf_info = tx_info.text_frame
    p_gv = tf_info.paragraphs[0]
    p_gv.text = f"Giáo viên thực hiện: {profile.get('full_name', '')}"
    p_gv.font.size = Pt(20)
    p_gv.font.bold = True
    p_gv.font.color.rgb = NAVY_COLOR

    p_cls = tf_info.add_paragraph()
    p_cls.text = f"Lớp giảng dạy: {slide_data.get('grade', profile.get('classes', ''))} • Năm học 2026 - 2027"
    p_cls.font.size = Pt(16)
    p_cls.font.color.rgb = TEXT_MAIN

    # Ghi chú cho slide 1
    s1.notes_slide.notes_text_frame.text = "Ghi chú GV: Chào cả lớp, ổn định trật tự và giới thiệu bài học mới."

    # 2. Các slide nội dung tiếp theo
    slides_list = slide_data.get('slides', [
        {
            "title": "MỤC TIÊU BÀI HỌC",
            "subtitle": "Yêu cầu cần đạt sau khi học xong bài này",
            "bullets": [
                "Nắm vững các định nghĩa, quy tắc và tính chất cơ bản của bài học.",
                "Biết cách áp dụng kiến thức vào giải quyết các bài tập thực hành.",
                "Phát triển năng lực tự chủ, giao tiếp và hợp tác nhóm tích cực."
            ],
            "notes": "Nhấn mạnh 3 mục tiêu trọng tâm để học sinh có định hướng học tập."
        },
        {
            "title": "HOẠT ĐỘNG 1: KHỞI ĐỘNG",
            "subtitle": "Kết nối kiến thức và tạo tình huống",
            "bullets": [
                "Quan sát hình ảnh thực tế trên máy chiếu.",
                "Thảo luận nhanh theo bàn (2 phút): Đâu là yếu tố quen thuộc?",
                "Đại diện 1-2 bạn học sinh nêu ý kiến trước lớp."
            ],
            "notes": "Dành khoảng 5 phút cho hoạt động khởi động, khích lệ tinh thần học sinh."
        },
        {
            "title": "HOẠT ĐỘNG 2: HÌNH THÀNH KIẾN THỨC",
            "subtitle": "Khám phá và xây dựng nội dung chính",
            "bullets": [
                "Nghiên cứu mục 1 trong Sách giáo khoa kết hợp thảo luận nhóm 4.",
                "Thực hiện phiếu học tập số 1 và ghi nhận kết quả.",
                "Rút ra kết luận khoa học và ghi vào vở bài học."
            ],
            "notes": "Theo dõi các nhóm thực hiện phiếu học tập, hỗ trợ học sinh gặp khó khăn."
        },
        {
            "title": "HOẠT ĐỘNG 3: LUYỆN TẬP & VẬN DỤNG",
            "subtitle": "Thực hành bài tập và mở rộng liên hệ",
            "bullets": [
                "Luyện tập bài tập 1, bài tập 2 theo mức độ nhận biết và thông hiểu.",
                "Vận dụng giải quyết tình huống thực tiễn mở rộng.",
                "Hướng dẫn học sinh nhiệm vụ tự học và chuẩn bị bài mới ở nhà."
            ],
            "notes": "Chốt lại các lỗi sai học sinh hay mắc phải và tổng kết tiết học."
        }
    ])

    for item in slides_list:
        slide = prs.slides.add_slide(blank_layout)

        # Header bar
        hbar = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(1.3))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = NAVY_COLOR
        hbar.line.color.rgb = NAVY_COLOR

        # Header Title Text
        htext = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.9))
        htf = htext.text_frame
        hp = htf.paragraphs[0]
        hp.text = item.get("title", "")
        hp.font.size = Pt(28)
        hp.font.bold = True
        hp.font.color.rgb = WHITE

        # Subtitle
        sub_shape = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.6))
        sub_tf = sub_shape.text_frame
        sub_p = sub_tf.paragraphs[0]
        sub_p.text = item.get("subtitle", "")
        sub_p.font.size = Pt(17)
        sub_p.font.italic = True
        sub_p.font.color.rgb = PRIMARY_COLOR

        # Content Box
        cbox = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(4.5))
        ctf = cbox.text_frame
        ctf.word_wrap = True

        bullets = item.get("bullets", [])
        for i, bullet in enumerate(bullets):
            p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            p.text = f"●   {bullet}"
            p.font.size = Pt(20)
            p.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(18)

        # Footer
        ft = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.5))
        ft_p = ft.text_frame.paragraphs[0]
        ft_p.text = f"{profile.get('school_name', '')} • Giáo viên: {profile.get('full_name', '')}"
        ft_p.font.size = Pt(11)
        ft_p.font.color.rgb = RGBColor(140, 150, 170)

        # Ghi chú thuyết minh (Speaker Notes)
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = item.get("notes")

    prs.save(output_path)
    return output_path
