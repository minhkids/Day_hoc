"""Independent preschool WPF app; shared document/AI primitives, isolated data."""
from __future__ import annotations
import csv
from contextlib import closing
import hashlib
import hmac
import json
import os
from pathlib import Path
import re
import secrets
import sqlite3
import subprocess
import sys
import tempfile
import threading
import unicodedata
import uuid
import zipfile
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

if not getattr(sys, 'frozen', False):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'specialist'))
import core
from catalog import CATALOG, PAGES, AGES, FIELDS, SYSTEM, THEMES, ADV_METHODS, PARENT_SCENARIOS, CREATIVE_GENRES, SITUATIONS
import legal
import ward
import assistant

try:
    import updater
except ImportError:
    try:
        from specialist import updater
    except ImportError:
        updater = None

ROOT = Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parent))
APP_ID = 'preschool'
APP_NAME = 'Trợ lý Giáo viên Mầm non'
VERSION = '1.0.4'
core.SYSTEM = SYSTEM
DEFAULTS = dict(name='', agency='', classes='', age=AGES[4], year='2026–2027',
                role='Giáo viên mầm non', position='Giáo viên mầm non',
                provider='OpenRouter', model='', endpoint='https://openrouter.ai/api/v1', provider_routing='')
core.DEFAULTS.update(DEFAULTS)
KINDS = {'document', 'observation', 'care', 'task', 'child'}

def atomic_json(path, data):
    temp = path.with_suffix('.tmp')
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    temp.replace(path)

class Store(core.Store):
    """SQLite holds IDs, types, dates and file paths; all content lives in files."""
    def profile(self):
        path = self.root / 'profile.json'
        return DEFAULTS | (json.loads(path.read_text(encoding='utf-8')) if path.exists() else {})

    def save_profile(self, data):
        profile = self.profile() | {k: str(v).strip() for k,v in data.items() if k in DEFAULTS}
        profile.update(provider='OpenRouter', endpoint=DEFAULTS['endpoint'], role=DEFAULTS['role'], position=DEFAULTS['position'])
        if profile['age'] not in AGES: raise ValueError('Chọn nhóm tuổi trong danh sách.')
        atomic_json(self.root / 'profile.json', profile)
        return profile

    def key(self):
        path = self.root / 'api-key.dpapi'
        return core.protect(path.read_text(encoding='ascii'), True) if path.exists() else ''

    def all(self, kind):
        records = super().all(kind)
        result = []
        for meta in records:
            path = (self.root / meta['path']).resolve()
            if not path.is_relative_to(self.root.resolve()): raise ValueError('Đường dẫn dữ liệu không hợp lệ.')
            if not path.is_file(): raise ValueError('Thiếu tệp dữ liệu: ' + meta['path'])
            result.append(json.loads(path.read_text(encoding='utf-8')))
        return result

    def put(self, kind, item):
        if kind not in KINDS: raise ValueError('Loại hồ sơ không hợp lệ.')
        item = dict(item)
        ident = item.get('id') or uuid.uuid4().hex
        if not isinstance(ident, str) or len(ident) != 32 or any(c not in '0123456789abcdef' for c in ident):
            raise ValueError('Mã hồ sơ không hợp lệ.')
        existing = self.db.execute('SELECT kind FROM records WHERE id=?', (ident,)).fetchone()
        if existing and existing[0] != kind: raise ValueError('Không thể đổi loại hồ sơ.')
        item.update(id=ident, updated=datetime.now().isoformat(timespec='seconds'))
        folder = self.root / 'ho-so' / kind
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / (ident + '.json')
        atomic_json(path, item)
        super().put(kind, dict(id=ident, path=str(path.relative_to(self.root)), updated=item['updated']))
        return item

    def delete(self, ident):
        row = self.db.execute('SELECT payload FROM records WHERE id=?',(ident,)).fetchone()
        if row:
            path = (self.root / json.loads(row[0])['path']).resolve()
            if not path.is_relative_to((self.root/'ho-so').resolve()): raise ValueError('Đường dẫn không hợp lệ.')
            # Preserve recoverable content outside the active index.
            trash = self.root / 'thung-rac'; trash.mkdir(exist_ok=True)
            if path.exists(): path.replace(trash / (uuid.uuid4().hex + '.json'))
            super().delete(ident)

    def backup(self, target):
        target = Path(target).resolve()
        if target.is_relative_to(self.root.resolve()): raise ValueError('Chọn nơi sao lưu bên ngoài thư mục dữ liệu.')
        with tempfile.TemporaryDirectory() as tmp:
            snapshot=Path(tmp)/'du-lieu.sqlite3'
            with closing(sqlite3.connect(snapshot)) as con: self.db.backup(con)
            with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
                z.write(snapshot,snapshot.name)
                for folder in ('ho-so','van-ban-da-soan','thung-rac'):
                    for path in (self.root/folder).rglob('*'):
                        if path.is_file(): z.write(path,path.relative_to(self.root))
                if (self.root/'profile.json').exists(): z.write(self.root/'profile.json','profile.json')

def export_document(path, title, body, fmt):
    if fmt == 'docx':
        from docx import Document
        from docx.shared import Pt, Cm
        doc=Document();doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(13)
        doc.sections[0].left_margin=Cm(2.5)
        doc.add_heading(title,0)
        doc.add_paragraph('DỰ THẢO — Giáo viên kiểm tra và điều chỉnh trước khi sử dụng.')
        for line in body.splitlines():
            if line.startswith('#'): doc.add_heading(line.lstrip('# ').strip(),min(len(line)-len(line.lstrip('#')),3))
            else: doc.add_paragraph(line)
        doc.save(path)
    elif fmt == 'pptx':
        core.export_slides(path,title,body)
        from pptx import Presentation
        prs=Presentation(path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs: run.text=run.text.replace('CHUYÊN VIÊN GIÁO DỤC','GIÁO VIÊN MẦM NON')
        prs.save(path)
    elif fmt == 'xlsx': core.export_table(path,[[line] for line in body.splitlines()],['Nội dung'])
    elif fmt == 'md': Path(path).write_text(body,encoding='utf-8')
    else: raise ValueError('Định dạng không hỗ trợ.')

def create_child_excel_template(target_path: str | Path) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Danh sách trẻ'

    # Title Banner
    ws.merge_cells('A1:Q1')
    t_cell = ws['A1']
    t_cell.value = 'DANH SÁCH HỌC SINH LỚP MẦM NON (FILE MẪU NHẬP LIỆU)'
    t_cell.font = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    t_cell.fill = PatternFill('solid', fgColor='164B43')
    t_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    # Subtitle / Guidance
    ws.merge_cells('A2:Q2')
    s_cell = ws['A2']
    s_cell.value = (
        'Hướng dẫn: Điền thông tin theo các cột bên dưới (cột có dấu * là bắt buộc). '
        'Nơi ở chỉ cần điền tên Phường/Xã cũ (kèm Quận/Huyện, Tỉnh/TP nếu có). '
        'Sau khi lưu, dùng nút [Nạp danh sách từ Excel] trên phần mềm để nạp vào lớp.'
    )
    s_cell.font = Font(name='Segoe UI', size=10, italic=True, color='164B43')
    s_cell.fill = PatternFill('solid', fgColor='E8F4F0')
    s_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 28

    headers = [
        'STT',
        'Mã trẻ',
        'Họ và tên trẻ (*)',
        'Giới tính',
        'Ngày sinh (yyyy-mm-dd)',
        'Nhóm / lớp',
        'Cân nặng (kg)',
        'Chiều cao (cm)',
        'Thể trạng dinh dưỡng',
        'Họ tên Bố',
        'SĐT Bố',
        'Họ tên Mẹ',
        'SĐT Mẹ',
        'Nơi ở trước sáp nhập (Phường/Xã cũ)',
        'Nơi ở sau sáp nhập (Phường/Xã mới)',
        'SĐT liên hệ khẩn cấp',
        'Ghi chú sức khỏe / Dị ứng'
    ]

    ws.append(headers)
    ws.row_dimensions[3].height = 28

    sample_rows = [
        [
            1, 'TRE-01', 'Nguyễn Minh Khôi', 'Nam', '2021-05-15', 'Lớp Mẫu giáo Lớn A1',
            16.5, 105.0, 'Bình thường (Kênh A)', 'Nguyễn Văn Nam', '0912345678',
            'Trần Thị Mai', '0987654321', 'Phường Phương Liên, Quận Đống Đa, Hà Nội',
            'Phường Kim Liên, Quận Đống Đa, Hà Nội', '0912345678', 'Không dị ứng'
        ],
        [
            2, 'TRE-02', 'Lê Bảo Ngọc', 'Nữ', '2021-08-20', 'Lớp Mẫu giáo Lớn A1',
            15.0, 102.0, 'Bình thường (Kênh A)', 'Lê Quốc Trung', '0903112233',
            'Vũ Hoàng Yến', '0903445566', 'Ngọc Thụy, Long Biên, Hà Nội',
            '', '0903112233', 'Cần nhắc uống thêm nước'
        ],
        [
            3, 'TRE-03', 'Trần Gia Hưng', 'Nam', '2021-03-10', 'Lớp Mẫu giáo Lớn A1',
            17.0, 106.0, 'Bình thường (Kênh A)', 'Trần Hải Nam', '0978123456',
            'Phạm Quỳnh Nga', '0978987654', 'Phường 5, Quận 3, TP.HCM',
            '', '0978123456', 'Dị ứng đậu phộng'
        ]
    ]

    for row in sample_rows:
        ws.append(row)

    thin = Side(border_style='thin', color='B4D8CB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Style Header row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='207465')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    # Style Data rows
    for r_idx in range(4, 4 + len(sample_rows)):
        ws.row_dimensions[r_idx].height = 24
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = Font(name='Segoe UI', size=10)
            cell.border = border
            if col_idx in (1, 2, 4, 5, 7, 8):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Column widths
    widths = [6, 12, 22, 11, 22, 22, 14, 14, 22, 20, 14, 20, 14, 34, 34, 18, 26]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'C4'
    Path(target_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(target_path)
    return str(target_path)

def import_children_from_file(file_path: str | Path, store: Store, root_dir: Path) -> dict:
    from openpyxl import load_workbook
    from datetime import datetime, date

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f'Tệp không tồn tại: {file_path}')

    rows_data: list[list[Any]] = []
    if path.suffix.lower() == '.csv':
        with open(path, 'r', encoding='utf-8-sig', errors='replace') as f:
            reader = csv.reader(f)
            rows_data = [row for row in reader]
    else:
        wb = load_workbook(filename=str(path), data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            rows_data.append(list(row))

    if not rows_data:
        raise ValueError('Tệp Excel không có dữ liệu.')

    header_row_idx = -1
    col_mapping: dict[str, int] = {}

    def norm_hdr(s: Any) -> str:
        if s is None:
            return ''
        s = str(s).strip().lower()
        return unicodedata.normalize('NFC', s)

    for r_idx, row in enumerate(rows_data[:10]):
        row_norm = [norm_hdr(c) for c in row]
        has_name = any(('họ' in c or 'tên' in c or 'name' in c) and 'bố' not in c and 'mẹ' not in c for c in row_norm)
        has_code = any('mã' in c or 'code' in c or 'stt' in c for c in row_norm)
        has_gender = any('giới tính' in c or 'gender' in c or 'phái' in c for c in row_norm)
        if has_name and (has_code or has_gender or len([c for c in row_norm if c]) >= 3):
            header_row_idx = r_idx
            for c_idx, cell_str in enumerate(row_norm):
                if not cell_str:
                    continue
                if any(k in cell_str for k in ('họ và tên', 'họ tên', 'tên trẻ', 'tên học sinh', 'tên', 'name')) and 'bố' not in cell_str and 'mẹ' not in cell_str:
                    if 'name' not in col_mapping: col_mapping['name'] = c_idx
                elif any(k in cell_str for k in ('mã trẻ', 'mã định danh', 'mã học sinh', 'số danh bộ', 'mã số', 'mã', 'code')):
                    if 'code' not in col_mapping: col_mapping['code'] = c_idx
                elif any(k in cell_str for k in ('giới tính', 'phái', 'gender', 'sex')):
                    if 'gender' not in col_mapping: col_mapping['gender'] = c_idx
                elif any(k in cell_str for k in ('ngày sinh', 'năm sinh', 'sinh nhật', 'dob', 'birth')):
                    if 'dob' not in col_mapping: col_mapping['dob'] = c_idx
                elif any(k in cell_str for k in ('nhóm / lớp', 'nhóm lớp', 'lớp', 'tên lớp', 'class')):
                    if 'class_name' not in col_mapping: col_mapping['class_name'] = c_idx
                elif any(k in cell_str for k in ('cân nặng', 'weight', 'cân nặng (kg)', 'kg')):
                    if 'weight' not in col_mapping: col_mapping['weight'] = c_idx
                elif any(k in cell_str for k in ('chiều cao', 'height', 'chiều cao (cm)', 'cm')):
                    if 'height' not in col_mapping: col_mapping['height'] = c_idx
                elif any(k in cell_str for k in ('thể trạng', 'dinh dưỡng', 'kênh', 'nutrition')):
                    if 'nutrition' not in col_mapping: col_mapping['nutrition'] = c_idx
                elif any(k in cell_str for k in ('sđt bố', 'số điện thoại bố', 'đt bố', 'điện thoại bố', 'phone bố', 'father phone')):
                    if 'father_phone' not in col_mapping: col_mapping['father_phone'] = c_idx
                elif any(k in cell_str for k in ('họ tên bố', 'tên bố', 'cha', 'bố', 'father')):
                    if 'father_name' not in col_mapping: col_mapping['father_name'] = c_idx
                elif any(k in cell_str for k in ('sđt mẹ', 'số điện thoại mẹ', 'đt mẹ', 'điện thoại mẹ', 'phone mẹ', 'mother phone')):
                    if 'mother_phone' not in col_mapping: col_mapping['mother_phone'] = c_idx
                elif any(k in cell_str for k in ('họ tên mẹ', 'tên mẹ', 'mẹ', 'mother')):
                    if 'mother_name' not in col_mapping: col_mapping['mother_name'] = c_idx
                elif any(k in cell_str for k in ('trước sáp nhập', 'địa chỉ cũ', 'nơi ở cũ', 'phường cũ', 'xã cũ', 'address_old', 'cũ')):
                    if 'address_old' not in col_mapping: col_mapping['address_old'] = c_idx
                elif any(k in cell_str for k in ('sau sáp nhập', 'địa chỉ mới', 'nơi ở mới', 'phường mới', 'xã mới', 'address_new', 'mới')):
                    if 'address_new' not in col_mapping: col_mapping['address_new'] = c_idx
                elif any(k in cell_str for k in ('địa chỉ', 'nơi ở', 'hộ khẩu', 'thường trú', 'tạm trú', 'address')):
                    if 'address_old' not in col_mapping: col_mapping['address_old'] = c_idx
                elif any(k in cell_str for k in ('sđt khẩn cấp', 'khẩn cấp', 'người giám hộ', 'liên hệ khẩn cấp', 'emergency')):
                    if 'emergency_phone' not in col_mapping: col_mapping['emergency_phone'] = c_idx
                elif any(k in cell_str for k in ('ghi chú', 'dị ứng', 'sức khỏe', 'lưu ý', 'health_notes', 'notes')):
                    if 'health_notes' not in col_mapping: col_mapping['health_notes'] = c_idx
            break

    if header_row_idx == -1 or 'name' not in col_mapping:
        raise ValueError('Không tìm thấy cột Họ và tên học sinh trong file Excel. Vui lòng tải file mẫu để xem định dạng chuẩn.')

    def get_cell_val(row: list[Any], key: str) -> str:
        idx = col_mapping.get(key)
        if idx is None or idx >= len(row):
            return ''
        v = row[idx]
        if v is None:
            return ''
        if isinstance(v, (datetime, date)):
            return v.strftime('%Y-%m-%d')
        return str(v).strip()

    imported: list[dict] = []
    skipped = 0
    current_count = len(store.all('child'))

    for r_idx in range(header_row_idx + 1, len(rows_data)):
        row = rows_data[r_idx]
        name = get_cell_val(row, 'name')
        if not name or name.lower() in ('họ và tên trẻ (*)', 'họ và tên', 'họ tên', 'nguyễn văn a', 'tên trẻ'):
            skipped += 1
            continue

        current_count += 1
        code = get_cell_val(row, 'code') or f'TRE-{str(current_count).zfill(2)}'
        gender_raw = get_cell_val(row, 'gender') or 'Nam'
        if any(w in gender_raw.lower() for w in ('nữ', 'nu', 'female', 'g')):
            gender = 'Nữ'
        else:
            gender = 'Nam'

        dob = get_cell_val(row, 'dob')
        if dob:
            m_vn = re.match(r'^(\d{1,2})[/\-](\d{1,2})[/\-](\d{4})$', dob)
            if m_vn:
                d, m, y = m_vn.groups()
                dob = f'{y}-{m.zfill(2)}-{d.zfill(2)}'
            else:
                m_iso = re.match(r'^(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})', dob)
                if m_iso:
                    y, m, d = m_iso.groups()
                    dob = f'{y}-{m.zfill(2)}-{d.zfill(2)}'

        weight_str = get_cell_val(row, 'weight').replace(',', '.')
        height_str = get_cell_val(row, 'height').replace(',', '.')
        try:
            weight = float(weight_str) if weight_str else ''
        except ValueError:
            weight = ''
        try:
            height = float(height_str) if height_str else ''
        except ValueError:
            height = ''

        growth = ward.evaluate_growth(weight, height) if (weight or height) else {'status': 'Bình thường (Kênh A)', 'bmi': ''}
        nutrition = get_cell_val(row, 'nutrition') or growth.get('status', 'Bình thường (Kênh A)')

        item = {
            'code': code,
            'name': name,
            'gender': gender,
            'dob': dob or datetime.now().strftime('%Y-%m-%d'),
            'class_name': get_cell_val(row, 'class_name') or store.profile().get('classes', ''),
            'weight': weight,
            'height': height,
            'nutrition': nutrition,
            'bmi': growth.get('bmi', ''),
            'father_name': get_cell_val(row, 'father_name'),
            'father_phone': get_cell_val(row, 'father_phone'),
            'mother_name': get_cell_val(row, 'mother_name'),
            'mother_phone': get_cell_val(row, 'mother_phone'),
            'emergency_phone': get_cell_val(row, 'emergency_phone') or get_cell_val(row, 'father_phone') or get_cell_val(row, 'mother_phone'),
            'address_old': get_cell_val(row, 'address_old'),
            'address_new': get_cell_val(row, 'address_new'),
            'health_notes': get_cell_val(row, 'health_notes')
        }

        saved = store.put('child', item)
        imported.append(saved)

    return {
        'count': len(imported),
        'skipped': skipped,
        'path': str(path),
        'total': len(store.all('child'))
    }

class Bridge:
    def __init__(self, root=None):
        self.root=Path(root or os.environ.get('TROLY_PRESCHOOL_DATA_DIR') or Path(os.environ.get('LOCALAPPDATA',str(Path.home())))/'TroLyGiaoVienMamNon').resolve()
        store=Store(self.root);store.db.close()
        self.jobs={};self.chats={};self.lock=threading.RLock()
        # Auto-update decrees & circulars in background on app start
        threading.Thread(target=self._auto_update_legal, daemon=True).start()

    def _auto_update_legal(self):
        try:
            legal.auto_update_legal_if_needed(self.root)
        except Exception:
            pass

    def dispatch(self, action, data):
        store=Store(self.root)
        try:
            if action=='state':
                fields=FIELDS if store.profile()['age'].startswith('Mẫu giáo') else FIELDS[:3]+['Tình cảm, kỹ năng xã hội và thẩm mỹ','Tích hợp theo độ tuổi']
                return dict(profile=store.profile(),has_key=bool(store.key()),root=str(self.root),version=VERSION,
                            catalog=CATALOG,pages=PAGES,ages=AGES,fields=fields,
                            themes=THEMES,adv_methods=ADV_METHODS,parent_scenarios=PARENT_SCENARIOS,
                            creative_genres=CREATIVE_GENRES,situations=SITUATIONS,
                            legal=legal.get_legal_docs(self.root),
                            legal_sync=legal.get_legal_sync_status(self.root),
                            ward_rules=ward.get_ward_rules(self.root),
                            **{k:store.all(k) for k in KINDS})
            if action=='ai_activity_wizard':
                profile, key = store.profile(), store.key()
                content = assistant.generate_activity_plan(
                    topic=data.get('topic', ''),
                    age=data.get('age') or profile.get('age', ''),
                    theme=data.get('theme', ''),
                    method=data.get('method', ''),
                    duration=data.get('duration', ''),
                    materials=data.get('materials', ''),
                    school=profile.get('agency', 'Trường Mầm non'),
                    class_name=profile.get('classes', 'Lớp Mầm non'),
                    profile=profile, key=key,
                    ask_ai_fn=core.ask_ai if (key and profile.get('model')) else None
                )
                return dict(content=content)
            if action=='ai_parent_message':
                profile, key = store.profile(), store.key()
                content = assistant.generate_parent_message(
                    scenario=data.get('scenario', ''),
                    child_name=data.get('child_name', ''),
                    details=data.get('details', ''),
                    class_name=profile.get('classes', 'Lớp Mầm non'),
                    teacher_name=profile.get('name', 'Cô giáo'),
                    school=profile.get('agency', 'Trường Mầm non'),
                    profile=profile, key=key,
                    ask_ai_fn=core.ask_ai if (key and profile.get('model')) else None
                )
                return dict(content=content)
            if action=='ai_refine_observation':
                profile, key = store.profile(), store.key()
                return assistant.refine_pedagogical_observation(
                    child=data.get('child', ''),
                    raw_notes=data.get('notes', ''),
                    area=data.get('area', ''),
                    kind=data.get('kind', 'observation'),
                    age=profile.get('age', 'Mẫu giáo 4–5 tuổi'),
                    profile=profile, key=key,
                    ask_ai_fn=core.ask_ai if (key and profile.get('model')) else None
                )
            if action=='ai_creative_studio':
                profile, key = store.profile(), store.key()
                content = assistant.generate_creative_content(
                    genre=data.get('genre', ''),
                    children_names=data.get('children_names', ''),
                    topic=data.get('topic', ''),
                    age=profile.get('age', 'Mẫu giáo 4–5 tuổi'),
                    profile=profile, key=key,
                    ask_ai_fn=core.ask_ai if (key and profile.get('model')) else None
                )
                return dict(content=content)
            if action=='ai_situation_help':
                profile, key = store.profile(), store.key()
                content = assistant.generate_situation_coaching(
                    situation=data.get('situation', ''),
                    details=data.get('details', ''),
                    age=profile.get('age', 'Mẫu giáo 4–5 tuổi'),
                    profile=profile, key=key,
                    ask_ai_fn=core.ask_ai if (key and profile.get('model')) else None
                )
                return dict(content=content)
            if action=='legal_sync':
                return legal.fetch_online_legal_updates(self.root)
            if action=='legal_verify':
                return legal.verify_teacher_legal_link(data.get('url',''), data.get('number',''), data.get('title',''))
            if action=='legal_verify_all':
                return legal.verify_all_docs(self.root)
            if action=='legal_repair_link':
                doc_id = data.get('id', '')
                canonical_url = legal.get_canonical_url(doc_id)
                if canonical_url:
                    docs = legal.get_legal_docs(self.root)
                    for d in docs:
                        if d.get('id') == doc_id:
                            d['url'] = canonical_url
                            d['verified'] = True
                            d['verification_note'] = 'Đã khôi phục đường dẫn chuẩn xác.'
                    legal.save_legal_docs(self.root, docs)
                    return dict(ok=True, url=canonical_url)
                return dict(ok=False, error='Không tìm thấy link chuẩn cho mã này.')
            if action=='ward_convert':
                direction = data.get('direction', 'old_to_new')
                address = data.get('address', '')
                api_key = data.get('api_key') or store.profile().get('geovina_api_key', '')
                if direction == 'new_to_old':
                    return ward.convert_address_new_to_old(address, self.root, api_key=api_key)
                return ward.convert_address_old_to_new(address, self.root, api_key=api_key)
            if action=='ward_batch_convert':
                addresses = data.get('addresses', [])
                api_key = data.get('api_key') or store.profile().get('geovina_api_key', '')
                return dict(results=ward.batch_convert_addresses(addresses, self.root, api_key=api_key))
            if action=='ward_rules':
                return dict(rules=ward.get_ward_rules(self.root))
            if action=='ward_save_rule':
                return dict(rules=ward.save_custom_ward_rule(self.root, data.get('rule', {})))
            if action=='evaluate_growth':
                return ward.evaluate_growth(data.get('weight'), data.get('height'))
            if action=='settings':
                profile=store.save_profile(data['profile'])
                if data.get('key'): store.save_key(data['key'])
                if data.get('remove_key'): store.save_key('')
                return profile
            if action=='save':
                kind=data['kind'];item=dict(data['item'])
                if kind not in KINDS: raise ValueError('Loại hồ sơ không hợp lệ.')
                if kind=='child':
                    if not str(item.get('name','')).strip(): raise ValueError('Vui lòng nhập họ và tên trẻ.')
                    if not str(item.get('code','')).strip():
                        item['code'] = 'TRE-' + str(len(store.all('child')) + 1).zfill(2)
                    growth = ward.evaluate_growth(item.get('weight'), item.get('height'))
                    if not item.get('nutrition'):
                        item['nutrition'] = growth['status']
                    item['bmi'] = growth['bmi']
                    if data.get('auto_convert') and item.get('address_old') and not item.get('address_new'):
                        geo_key = store.profile().get('geovina_api_key', '')
                        conv = ward.convert_address_old_to_new(item['address_old'], self.root, api_key=geo_key)
                        if conv.get('converted'):
                            item['address_new'] = conv['result']
                if kind in ('observation','care'):
                    if not str(item.get('child','')).strip(): raise ValueError('Nhập mã hoặc tên trẻ / nhóm.')
                    datetime.strptime(item.get('date',''),'%Y-%m-%d')
                    if not str(item.get('notes','')).strip(): raise ValueError('Nhập ghi nhận thực tế.')
                if kind=='task':
                    if not str(item.get('title','')).strip(): raise ValueError('Nhập tên công việc.')
                    datetime.strptime(item.get('due',''),'%Y-%m-%d')
                if kind=='document' and (not str(item.get('title','')).strip() or not str(item.get('body','')).strip()):
                    raise ValueError('Nhập tiêu đề và nội dung tài liệu.')
                return store.put(kind,item)
            if action=='delete': store.delete(data['id']);return True
            if action=='export':
                path=Path(data['path']);export_document(path,data['title'],data['body'],data['format']);return dict(path=str(path))
            if action=='export_child_template':
                path=create_child_excel_template(data['path'])
                return dict(path=str(path))
            if action=='import_children':
                return import_children_from_file(data['path'], store, self.root)
            if action=='export_records':
                kind=data['kind']
                if kind not in ('observation','care','task','child'): raise ValueError('Loại sổ không hợp lệ.')
                rows=store.all(kind)
                if data.get('query'): rows=[r for r in rows if data['query'].casefold() in json.dumps(r,ensure_ascii=False).casefold()]
                if kind=='child':
                    cols=['code','name','gender','dob','class_name','weight','height','nutrition','father_name','father_phone','mother_name','mother_phone','address_old','address_new','emergency_phone','health_notes']
                    headers=['Mã trẻ','Họ và tên','Giới tính','Ngày sinh','Nhóm / lớp','Cân nặng (kg)','Chiều cao (cm)','Thể trạng','Họ tên Bố','SĐT Bố','Họ tên Mẹ','SĐT Mẹ','Nơi ở (Trước sáp nhập)','Nơi ở (Sau sáp nhập)','SĐT khẩn cấp','Ghi chú sức khỏe']
                elif kind=='observation':
                    cols=['date','child','area','status','notes','next']
                    headers=['Ngày','Trẻ / nhóm','Lĩnh vực','Mức ghi nhận','Quan sát thực tế','Hỗ trợ tiếp theo']
                elif kind=='care':
                    cols=['date','child','attendance','meal','sleep','notes','next']
                    headers=['Ngày','Trẻ / nhóm','Chuyên cần','Ăn / vệ sinh','Ngủ','Ghi nhận','Trao đổi gia đình']
                else:
                    cols=['title','due','status']
                    headers=['Công việc','Hạn','Trạng thái']
                core.export_table(data['path'],[[r.get(c,'') for c in cols] for r in rows],headers)
                return dict(path=data['path'],count=len(rows))
            if action=='read': return dict(text=core.read_document(data['path']))
            if action=='backup': store.backup(data['path']);return dict(path=data['path'])
            if action=='chat_clear':
                with self.lock:
                    self.chats.clear()
                    for job in self.jobs.values():
                        if job['state']=='running': job['state']='cancelled'
                return True
            if action=='chat_start':
                prompt=str(data.get('prompt','')).strip()
                if not prompt: raise ValueError('Nhập yêu cầu trước khi gửi.')
                profile,key=store.profile(),store.key()
                if not profile['model'] or not key: raise ValueError('Vào Cài đặt để nhập model và API key OpenRouter. Các mẫu và sổ theo dõi vẫn dùng được khi chưa có AI.')
                with self.lock:
                    if any(j['state']=='running' for j in self.jobs.values()): raise ValueError('Hãy chờ yêu cầu đang xử lý hoặc hủy trước.')
                    sid=data.get('session') or uuid.uuid4().hex;job=uuid.uuid4().hex
                    history=list(self.chats.get(sid,[]));self.jobs[job]=dict(state='running',session=sid)
                threading.Thread(target=self.generate,args=(job,sid,profile,key,history,prompt,data.get('files',[])),daemon=True).start()
                return dict(job=job,session=sid)
            if action=='chat_poll':
                with self.lock: return dict(self.jobs.get(data['job'],dict(state='missing')))
            if action=='chat_cancel':
                with self.lock:
                    job=self.jobs.get(data['job'])
                    if job and job['state']=='running': job['state']='cancelled'
                return True
            if action=='check_update':
                if updater:
                    return updater.check_for_updates(
                        APP_ID, VERSION,
                        interactive=data.get('interactive', True),
                        app_title=APP_NAME
                    )
                return {'status': 'error', 'error': 'Chưa kích hoạt module cập nhật.'}
            raise ValueError('Thao tác không hỗ trợ.')
        finally: store.db.close()

    def generate(self, job, sid, profile, key, history, prompt, files):
        try:
            content=prompt
            for file in files:
                content+='\n<TAI_LIEU>\n'+core.read_document(file)+'\n</TAI_LIEU>'
            answer=core.ask_ai(profile,key,history+[dict(role='user',content=content)])
            with self.lock:
                if self.jobs[job]['state']!='running': return
                self.chats[sid]=history+[dict(role='user',content=content),dict(role='assistant',content=answer)]
                self.jobs[job].update(state='done',answer=answer)
        except Exception as e:
            with self.lock:
                if self.jobs[job]['state']=='running': self.jobs[job].update(state='error',error=str(e))

def create_server(bridge, token):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,*args): pass
        def do_POST(self):
            if not hmac.compare_digest(self.headers.get('X-TroLy-Token',''),token): self.send_error(403);return
            try:
                length=int(self.headers.get('Content-Length','0'))
                if not 0<=length<=8*1024*1024: self.send_error(413);return
                result=dict(ok=True,data=bridge.dispatch(self.path.strip('/'),json.loads(self.rfile.read(length) or b'{}')))
            except Exception as e: result=dict(ok=False,error=str(e))
            data=json.dumps(result,ensure_ascii=False).encode('utf-8')
            self.send_response(200);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
    return ThreadingHTTPServer(('127.0.0.1',0),Handler)

def launch():
    if sys.platform == 'win32':
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID('TroLyGiaoDuc.MamNon.Desktop')
        except Exception:
            pass
    bridge=Bridge();token=secrets.token_urlsafe(32);server=create_server(bridge,token)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    if updater and '--smoke' not in sys.argv and '--screenshots' not in sys.argv:
        updater.start_background_check(APP_ID, VERSION, APP_NAME, delay_seconds=3.0)
    env=os.environ.copy();env.update(TROLY_BRIDGE_URL=f'http://127.0.0.1:{server.server_port}',TROLY_BRIDGE_TOKEN=token,TROLY_PRESCHOOL_DATA_DIR=str(bridge.root))
    args=[str(Path(os.environ.get('WINDIR','C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'),'-NoProfile','-STA','-ExecutionPolicy','Bypass','-WindowStyle','Hidden','-File',str(ROOT/'wpf/host.ps1')]
    if '--smoke' in sys.argv: args+=['-Smoke']
    if '--smoke-ai' in sys.argv and '--smoke' in sys.argv: args+=['-SmokeAI']
    if '--screenshots' in sys.argv: args+=['-ScreenshotDir',str(Path(sys.argv[sys.argv.index('--screenshots')+1]).resolve())]
    try: return subprocess.call(args,env=env,creationflags=subprocess.CREATE_NO_WINDOW)
    finally: server.shutdown();server.server_close()

if __name__=='__main__': sys.exit(launch())
