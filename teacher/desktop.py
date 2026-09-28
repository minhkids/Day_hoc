"""Teacher WPF desktop using the installed reference's presentation assets."""
from pathlib import Path
import os
import sys
import json
import threading
import uuid
import shutil
import zipfile
import re
import sqlite3
import tempfile
import secrets
import subprocess

if not getattr(sys,'frozen',False):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    sys.path.append(str(Path(__file__).resolve().parents[1]/'specialist'))

import core
import importlib.util
# Import the shared bridge under a distinct name (this module is also desktop.py).
if getattr(sys,'frozen',False):
    import shared_bridge as shared
else:
    spec=importlib.util.spec_from_file_location('shared_bridge',Path(__file__).resolve().parents[1]/'specialist/desktop.py')
    shared=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(shared)

if getattr(sys, 'frozen', False):
    from main import export_lesson, FIELDS  # Teacher prompt and defaults are initialized here.
else:
    spec_main = importlib.util.spec_from_file_location('teacher_main', Path(__file__).resolve().parent / 'main.py')
    teacher_main = importlib.util.module_from_spec(spec_main)
    spec_main.loader.exec_module(teacher_main)
    export_lesson, FIELDS = teacher_main.export_lesson, teacher_main.FIELDS



try:
    import updater
except ImportError:
    try:
        from specialist import updater
    except ImportError:
        updater = None

APP_ID = "teacher"
APP_NAME = "Trợ lý Giáo viên"
APP_VERSION = "1.0.7"

ROOT=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parent))

ASSETS=ROOT if getattr(sys,'frozen',False) else ROOT.parent/'app'
conversations={}


class Store(core.Store):
    def all(self,kind):
        if kind=='chat':return list(conversations.values())
        result=super().all(kind)
        if kind in ('document','template','memory'):
            for item in result:
                if item.get('path'):
                    path=(self.root/item['path']).resolve()
                    if path.is_relative_to(self.root.resolve()) and path.is_file():
                        item['body']=path.read_text(encoding='utf-8-sig')
        return result

    def put(self,kind,value):
        value=dict(value);value.setdefault('id',uuid.uuid4().hex)
        if kind=='chat':conversations[value['id']]=value;return value
        if kind in ('document','template','memory') and 'body' in value:
            folder={'document':'bai-day','template':'mau-rieng','memory':'bo-nho'}[kind]
            (self.root/folder).mkdir(exist_ok=True)
            path=self.root/folder/(core.safe_name(value['id'])+'.md')
            path.write_text(value.pop('body'),encoding='utf-8')
            value['path']=str(path.relative_to(self.root))
        return super().put(kind,value)

    def delete(self,key):
        conversations.pop(key,None);super().delete(key)

    def backup(self,target):
        with tempfile.TemporaryDirectory() as temp:
            snapshot=Path(temp)/'du-lieu.sqlite3'; con=sqlite3.connect(snapshot)
            try:self.db.backup(con)
            finally:con.close()
            with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
                z.write(snapshot,'du-lieu.sqlite3')
                for folder in ('bai-day','kiem-tra','chu-nhiem','ho-so','phan-mem','mau-rieng','bo-nho','van-ban-den','van-ban-da-soan'):
                    for p in (self.root/folder).rglob('*'):
                        if p.is_file() and p.resolve()!=Path(target).resolve():z.write(p,p.relative_to(self.root))


shared.Store=Store
shared.export_word=lambda path,title,body,profile:export_lesson(path,title,body)
def export_slides(path,title,body):
    core.export_slides(path,title,body)
    from pptx import Presentation
    presentation=Presentation(path)
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:run.text=run.text.replace('CHUYÊN VIÊN GIÁO DỤC','GIÁO VIÊN')
    presentation.save(path)
shared.export_slides=export_slides


class Bridge(shared.Bridge):
    def __init__(self,root=None):
        self.root=Path(root or os.environ.get('TROLY_TEACHER_DATA_DIR') or Path(os.environ.get('LOCALAPPDATA',str(Path.home())))/'TroLyGiaoVienDocLap')
        store=Store(self.root)
        for folder in ('bai-day','phan-mem','bo-nho'): (self.root/folder).mkdir(exist_ok=True)
        if not store.all('teacher_seed'):
            for p in sorted((ASSETS/'templates/giao_vien').glob('*.md')):
                store.put('template',{'title':p.stem,'body':p.read_text(encoding='utf-8-sig')})
            store.put('teacher_seed',{'id':'teacher-templates-v1'})
        store.db.close();self.jobs={};self.lock=threading.RLock()

    def dispatch(self,action,data):
        if action=='projects':
            return [{'title':p.name,'path':str(p)} for p in sorted((self.root/'phan-mem').iterdir()) if p.is_dir()]
        if action=='project_create':
            kind=data['kind']
            if kind not in ('trac-nghiem-doi','flashcard','vong-quay','dong-ho','o-chu','trong'):raise ValueError('Mẫu không hợp lệ')
            path=core.unique_path(self.root/'phan-mem',kind);shutil.copytree(ASSETS/'mini_apps'/kind,path)
            return {'path':str(path)}
        if action=='project_zip':
            path=Path(data['path']).resolve()
            if not path.is_relative_to((self.root/'phan-mem').resolve()):raise ValueError('Dự án không hợp lệ')
            with zipfile.ZipFile(data['target'],'w',zipfile.ZIP_DEFLATED) as z:
                for p in path.rglob('*'):
                    if p.is_file() and p.resolve()!=Path(data['target']).resolve():z.write(p,p.relative_to(path))
        if action == 'grade_stats':
            source = Path(data['path']).resolve()
            if not source.is_file():
                raise ValueError('Không tìm thấy tệp bảng điểm.')
            numbers = []
            if source.suffix.lower() == '.csv':
                import csv
                with source.open('r', encoding='utf-8-sig', errors='replace', newline='') as fh:
                    for row in csv.reader(fh):
                        for value in row:
                            try:
                                number = float(value.strip().replace(',', '.'))
                            except (ValueError, AttributeError):
                                continue
                            if 0 <= number <= 10:
                                numbers.append(number)
            else:
                text = core.read_document(source)
                numbers = [float(v.replace(',', '.')) for v in re.findall(r'(?<![\d.])(?:10(?:[.,]0+)?|[0-9](?:[.,][0-9]+)?)(?![\d.])', text)]
            if not numbers:
                raise ValueError('Không tìm thấy điểm số từ 0 đến 10 trong tệp.')
            passed = sum(value >= 5 for value in numbers)
            return {'count': len(numbers), 'average': round(sum(numbers) / len(numbers), 2),
                    'minimum': min(numbers), 'maximum': max(numbers), 'passed': passed,
                    'failed': len(numbers) - passed, 'pass_rate': round(passed * 100 / len(numbers), 1)}
        if action == 'check_update':
            if updater:
                return updater.check_for_updates(
                    APP_ID, APP_VERSION,
                    interactive=data.get('interactive', True),
                    app_title=APP_NAME
                )
            return {'status': 'error', 'error': 'Chưa kích hoạt module cập nhật.'}
        return super().dispatch(action,data)


def launch():
    if '--verify-ai' in sys.argv and '--smoke' not in sys.argv:
        sys.argv.append('--smoke')
    bridge=Bridge();token=secrets.token_urlsafe(32);server=shared.create_server(bridge,token)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    
    # Tự động kiểm tra cập nhật từ xa trong luồng phụ khi khởi động
    if updater and '--smoke' not in sys.argv and '--screenshots' not in sys.argv:
        updater.start_background_check(APP_ID, APP_VERSION, APP_NAME, delay_seconds=3.0)

    env=os.environ.copy();env.update(TROLY_BRIDGE_URL=f'http://127.0.0.1:{server.server_port}',TROLY_BRIDGE_TOKEN=token)
    args=[str(Path(os.environ.get('WINDIR','C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'),'-NoProfile','-STA','-ExecutionPolicy','RemoteSigned','-WindowStyle','Hidden','-File',str(ROOT/'wpf/host.ps1')]
    if '--smoke' in sys.argv:args+=['-Smoke']
    if '--verify-ai' in sys.argv:args+=['-VerifyAI']
    if '--screenshots' in sys.argv:args+=['-ScreenshotDir',str(Path(sys.argv[sys.argv.index('--screenshots')+1]).resolve())]
    try:return subprocess.call(args,env=env,creationflags=subprocess.CREATE_NO_WINDOW)
    finally:server.shutdown();server.server_close()


if __name__=='__main__':sys.exit(launch())
