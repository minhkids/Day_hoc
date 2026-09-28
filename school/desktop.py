"""School management WPF desktop using the installed reference's presentation assets."""
from pathlib import Path
import os
import sys
import json
import threading
import uuid
import shutil
import zipfile
import sqlite3
import tempfile
import secrets
import subprocess
from datetime import datetime

if not getattr(sys, 'frozen', False):
    sys.path.append(str(Path(__file__).resolve().parents[1] / 'specialist'))
import core
import importlib.util
# Import the shared bridge under a distinct name (this module is also desktop.py).
if getattr(sys, 'frozen', False):
    import shared_bridge as shared
else:
    spec = importlib.util.spec_from_file_location('shared_bridge', Path(__file__).resolve().parents[1] / 'specialist/desktop.py')
    shared = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(shared)

try:
    import updater
except ImportError:
    try:
        from specialist import updater
    except ImportError:
        updater = None

APP_ID = "school"
APP_NAME = "Trợ lý Quản trị trường học"
APP_VERSION = "1.0.7"


core.DEFAULTS.update(
    role='Hiệu trưởng',
    position='Hiệu trưởng',
    school_level='TH',
    principal='Trần Thanh Chung',
    vice_principals='',
    party_org='Chi bộ Trường Tiểu học',
    school_team='',
    provider='OpenRouter',
    model='nvidia/nemotron-3-super-120b-a12b:free',
    provider_routing='nvidia',
    endpoint='https://openrouter.ai/api/v1'
)

core.SYSTEM = """Bạn là trợ lý quản trị trường học chuyên nghiệp, hỗ trợ đắc lực cho Hiệu trưởng, Phó Hiệu trưởng và Tổ trưởng chuyên môn.
Trả lời hoàn toàn bằng tiếng Việt chuẩn mực, căn cứ trên hồ sơ nhà trường, quy chế chuyên môn và quy định hiện hành của Bộ GD&ĐT.
Thể thức văn bản hành chính thực hiện đúng Nghị định 30/2020/NĐ-CP (Quốc hiệu, Tiêu ngữ, Tên cơ quan ban hành, Số/Ký hiệu, Trích yếu, Căn cứ ban hành, Nội dung điều khoản, Nơi nhận, Chức danh người ký).
Với văn bản của Đảng, tuân thủ chặt chẽ Hướng dẫn 05-HD/VPTW và Điều lệ Đảng Cộng sản Việt Nam.
Kế hoạch phải có mục tiêu, chỉ tiêu cụ thể, phân công người phụ trách, thời hạn và sản phẩm nghiệm thu.
Báo cáo phải nêu rõ số liệu, đánh giá ưu điểm, tồn tại, nguyên nhân và bài học kinh nghiệm.
Không bịa đặt số liệu hay hiệu lực văn bản. Đánh dấu [CẦN BỔ SUNG] nếu thiếu thông tin địa phương."""

ROOT = Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parent))
ASSETS = ROOT if getattr(sys, 'frozen', False) else ROOT.parent / 'app'
conversations = {}


class Store(core.Store):
    def all(self, kind):
        if kind == 'chat':
            return list(conversations.values())
        result = super().all(kind)
        if kind in ('document', 'template', 'memory'):
            for item in result:
                if item.get('path'):
                    path = (self.root / item['path']).resolve()
                    if path.is_relative_to(self.root.resolve()) and path.is_file():
                        item['body'] = path.read_text(encoding='utf-8-sig')
        return result

    def put(self, kind, value):
        value = dict(value)
        value.setdefault('id', uuid.uuid4().hex)
        if kind == 'chat':
            conversations[value['id']] = value
            return value
        if kind in ('document', 'template', 'memory') and 'body' in value:
            folder = {'document': 'van-ban-da-soan', 'template': 'mau-rieng', 'memory': 'bo-nho'}[kind]
            (self.root / folder).mkdir(exist_ok=True)
            path = self.root / folder / (core.safe_name(value['id']) + '.md')
            path.write_text(value.pop('body'), encoding='utf-8')
            value['path'] = str(path.relative_to(self.root))
        return super().put(kind, value)

    def delete(self, key):
        conversations.pop(key, None)
        super().delete(key)

    def backup(self, target):
        with tempfile.TemporaryDirectory() as temp:
            snapshot = Path(temp) / 'du-lieu.sqlite3'
            con = sqlite3.connect(snapshot)
            try:
                self.db.backup(con)
            finally:
                con.close()
            with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as z:
                z.write(snapshot, 'du-lieu.sqlite3')
                for folder in ('van-ban-da-soan', 'kiem-tra', 'chu-nhiem', 'ho-so', 'phan-mem', 'mau-rieng', 'bo-nho', 'van-ban-den'):
                    for p in (self.root / folder).rglob('*'):
                        if p.is_file() and p.resolve() != Path(target).resolve():
                            z.write(p, p.relative_to(self.root))


shared.Store = Store
shared.export_word = core.export_word


def export_slides(path, title, body):
    core.export_slides(path, title, body)
    try:
        from pptx import Presentation
        presentation = Presentation(path)
        for shape in presentation.slides[0].shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.text = run.text.replace('CHUYÊN VIÊN GIÁO DỤC', 'QUẢN TRỊ TRƯỜNG HỌC')
        presentation.save(path)
    except Exception:
        pass


shared.export_slides = export_slides


class Bridge(shared.Bridge):
    def __init__(self, root=None):
        self.root = Path(root or os.environ.get('TROLY_SCHOOL_DATA_DIR') or Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'TroLyQuanTriTruongHocDocLap')
        store = Store(self.root)
        for folder in ('van-ban-da-soan', 'phan-mem', 'bo-nho', 'van-ban-den', 'mau-rieng'):
            (self.root / folder).mkdir(exist_ok=True)

        # Seed templates
        if not store.all('school_seed'):
            templates_dir = ASSETS / 'templates/hieu_truong'
            if templates_dir.is_dir():
                for p in sorted(templates_dir.glob('*.md')):
                    store.put('template', {'title': p.stem, 'body': p.read_text(encoding='utf-8-sig')})
            store.put('school_seed', {'id': 'school-templates-v1'})

        # Tự động nạp cấu hình OpenRouter từ .env nếu chưa có
        dotenv = core._dotenv_values()
        if not store.key() and dotenv.get('OPENROUTER_KEY'):
            store.save_key(dotenv['OPENROUTER_KEY'])
        prof = store.profile()
        if not prof.get('model') and dotenv.get('MODEL'):
            prof['model'] = dotenv['MODEL']
            prof['provider'] = 'OpenRouter'
            prof['endpoint'] = 'https://openrouter.ai/api/v1'
            store.save_profile(prof)

        store.db.close()
        self.jobs = {}
        self.lock = threading.RLock()

    def dispatch(self, action, data):
        if action == 'missions':
            missions_path = ROOT / 'wpf/missions.json'
            if missions_path.exists():
                return json.loads(missions_path.read_text(encoding='utf-8'))
            return []

        if action == 'mission':
            missions = self.dispatch('missions', {})
            item = next((m for m in missions if m['id'] == data.get('id')), None)
            if not item:
                raise ValueError('Nhiệm vụ không hợp lệ')
            card = ASSETS / 'templates/hieu_truong/kien-thuc/nhiem-vu' / (item['id'] + '.md')
            return item | {'body': card.read_text(encoding='utf-8-sig') if card.exists() else item.get('description', '')}

        if action == 'templates':
            store = Store(self.root)
            try:
                return store.all('template')
            finally:
                store.db.close()

        if action == 'chats':
            store = Store(self.root)
            try:
                return store.all('chat')
            finally:
                store.db.close()

        if action == 'export_word':
            path = Path(data['path'])
            title = data.get('title', 'Van-ban')
            body = data.get('body', '')
            store = Store(self.root)
            try:
                core.export_word(path, title, body, store.profile())
                return {'path': str(path)}
            finally:
                store.db.close()

        if action == 'chat':
            prompt = data.get('prompt', '').strip()
            if not prompt:
                raise ValueError('Chưa nhập câu hỏi hoặc yêu cầu.')
            sid = data.get('id')
            store = Store(self.root)
            try:
                profile = store.profile()
                key = store.key()
                dotenv = core._dotenv_values()
                if not key and dotenv.get('OPENROUTER_KEY'):
                    key = dotenv['OPENROUTER_KEY']
                    store.save_key(key)
                if not profile.get('model'):
                    profile['model'] = dotenv.get('MODEL') or 'nvidia/nemotron-3-super-120b-a12b:free'
                    profile['provider'] = 'OpenRouter'
                    profile['provider_routing'] = dotenv.get('PROVIDER') or 'nvidia'
                    profile['endpoint'] = 'https://openrouter.ai/api/v1'
                elif dotenv.get('MODEL'):
                    profile['model'] = dotenv['MODEL']
                    profile['provider_routing'] = dotenv.get('PROVIDER') or 'nvidia'
                store.save_profile(profile)

                conversation = next((c for c in store.all('chat') if c['id'] == sid), None) if sid else None
                if not conversation:
                    conversation = store.put('chat', {'title': prompt[:65], 'created': datetime.now().isoformat(timespec='seconds'), 'messages': []})

                files = data.get('attachments') or []
                content = prompt
                for f in files:
                    p = Path(f.get('path', ''))
                    if p.is_file():
                        try:
                            content += f"\n\n<TAI_LIEU ten={json.dumps(p.name, ensure_ascii=False)}>\n{core.read_document(p)}\n</TAI_LIEU>"
                        except Exception:
                            pass

                messages = [{'role': m['role'], 'content': m['content']} for m in conversation['messages']]
                memory = '\n\n'.join(m['title'] + ':\n' + m['body'] for m in store.all('memory'))
                answer = core.ask_ai(profile, key, messages + [{'role': 'user', 'content': content}], memory)

                display = prompt + ('\nĐính kèm: ' + ', '.join(f.get('name', '') for f in files) if files else '')
                conversation['messages'] += [{'role': 'user', 'content': content, 'display': display}, {'role': 'assistant', 'content': answer}]
                store.put('chat', conversation)
                return conversation
            finally:
                store.db.close()

        if action == 'delete':
            key = data.get('key') or data.get('id')
            if key:
                store = Store(self.root)
                try:
                    store.delete(key)
                    return True
                finally:
                    store.db.close()

        if action == 'check_update':
            if updater:
                return updater.check_for_updates(
                    APP_ID, APP_VERSION,
                    interactive=data.get('interactive', True),
                    app_title=APP_NAME
                )
            return {'status': 'error', 'error': 'Chưa kích hoạt module cập nhật.'}

        return super().dispatch(action, data)


def launch():
    if '--verify-ai' in sys.argv and '--smoke' not in sys.argv:
        sys.argv.append('--smoke')
    bridge = Bridge()
    token = secrets.token_urlsafe(32)
    server = shared.create_server(bridge, token)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    # Tự động kiểm tra cập nhật từ xa trong luồng phụ khi khởi động
    if updater and '--smoke' not in sys.argv and '--screenshots' not in sys.argv:
        updater.start_background_check(APP_ID, APP_VERSION, APP_NAME, delay_seconds=3.0)

    env = os.environ.copy()

    env.update(TROLY_BRIDGE_URL=f'http://127.0.0.1:{server.server_port}', TROLY_BRIDGE_TOKEN=token)
    args = [
        str(Path(os.environ.get('WINDIR', 'C:/Windows')) / 'System32/WindowsPowerShell/v1.0/powershell.exe'),
        '-NoProfile', '-STA', '-ExecutionPolicy', 'RemoteSigned', '-WindowStyle', 'Hidden',
        '-File', str(ROOT / 'wpf/host.ps1')
    ]
    if '--smoke' in sys.argv:
        args += ['-Smoke']
    if '--verify-ai' in sys.argv:
        args += ['-VerifyAI']
    if '--screenshots' in sys.argv:
        args += ['-ScreenshotDir', str(Path(sys.argv[sys.argv.index('--screenshots') + 1]).resolve())]
    try:
        return subprocess.call(args, env=env, creationflags=subprocess.CREATE_NO_WINDOW)
    finally:
        server.shutdown()
        server.server_close()


if __name__ == '__main__':
    sys.exit(launch())
